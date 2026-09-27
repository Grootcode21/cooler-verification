import openpyxl
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from inventory.models import Unit


class Command(BaseCommand):
    help = (
        "Import units from the Excel master list. "
        "Expected columns: Id, cooler tag, serial number, asset number, "
        "location, status, last-verified"
    )

    def add_arguments(self, parser):
        parser.add_argument(
            'excel_file',
            type=str,
            help='Path to the .xlsx file'
        )
        parser.add_argument(
            '--dry-run',
            action='store_true',
            help='Validate and report without saving anything'
        )
        parser.add_argument(
            '--update',
            action='store_true',
            help='Update existing units instead of skipping them'
        )

    def handle(self, *args, **options):
        path = options['excel_file']
        dry_run = options['dry_run']
        do_update = options['update']

        # --- Load workbook ---
        try:
            wb = openpyxl.load_workbook(path, data_only=True)
        except FileNotFoundError:
            raise CommandError(f"File not found: {path}")
        except Exception as e:
            raise CommandError(f"Could not open workbook: {e}")

        ws = wb.active
        if ws.max_row < 2:
            raise CommandError("Workbook appears empty (no data rows).")

        # --- Read and validate headers ---
        raw_headers = [str(c.value).strip().lower() if c.value else '' for c in ws[1]]
        expected = [
            'id', 'cooler tag', 'serial number', 'asset number',
            'location', 'status', 'last-verified',
        ]

        missing = [h for h in expected if h not in raw_headers]
        if missing:
            raise CommandError(
                f"Missing expected column(s): {missing}\n"
                f"Found headers: {raw_headers}"
            )

        idx = {name: raw_headers.index(name) for name in expected}

        # --- Track counters and errors ---
        created = updated = skipped_dupe = skipped_invalid = 0
        errors = []

        # --- Process rows ---
        with transaction.atomic():
            for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                # Skip completely empty rows
                if not any(row):
                    continue

                tag = self._clean(row[idx['cooler tag']])
                serial = self._clean(row[idx['serial number']])
                asset = self._clean(row[idx['asset number']])
                location = self._clean(row[idx['location']])
                status = self._clean(row[idx['status']]) or 'Active'

                # --- Required fields ---
                if not tag or not serial or not asset:
                    errors.append(
                        f"Row {row_num}: missing required field(s) "
                        f"(tag='{tag}', serial='{serial}', asset='{asset}')"
                    )
                    skipped_invalid += 1
                    continue

                # --- Status sanity ---
                if status not in ['Active', 'Retired', 'Lost']:
                    errors.append(
                        f"Row {row_num}: unknown status '{status}', defaulting to Active"
                    )
                    status = 'Active'

                # --- Look for existing record by cooler_tag ---
                existing = Unit.objects.filter(cooler_tag=tag).first()

                if existing:
                    if do_update:
                        # Guard: ensure the new serial/asset aren't already used by a different unit
                        conflict_serial = Unit.objects.filter(
                            serial_number=serial
                        ).exclude(pk=existing.pk).exists()
                        conflict_asset = Unit.objects.filter(
                            asset_number=asset
                        ).exclude(pk=existing.pk).exists()

                        if conflict_serial or conflict_asset:
                            errors.append(
                                f"Row {row_num}: update would collide with another unit "
                                f"(serial conflict={conflict_serial}, asset conflict={conflict_asset})"
                            )
                            skipped_invalid += 1
                            continue

                        existing.serial_number = serial
                        existing.asset_number = asset
                        existing.location = location
                        existing.status = status
                        existing.save()
                        updated += 1
                    else:
                        skipped_dupe += 1
                    continue

                # --- New unit: check serial/asset uniqueness before insert ---
                if Unit.objects.filter(serial_number=serial).exists():
                    errors.append(f"Row {row_num}: serial '{serial}' already exists on another unit")
                    skipped_invalid += 1
                    continue

                if Unit.objects.filter(asset_number=asset).exists():
                    errors.append(f"Row {row_num}: asset '{asset}' already exists on another unit")
                    skipped_invalid += 1
                    continue

                Unit.objects.create(
                    cooler_tag=tag,
                    serial_number=serial,
                    asset_number=asset,
                    location=location,
                    status=status,
                )
                created += 1

            if dry_run:
                transaction.set_rollback(True)

        # --- Report ---
        prefix = "[DRY RUN] " if dry_run else ""
        self.stdout.write(self.style.SUCCESS(
            f"\n{prefix}Import complete:"
        ))
        self.stdout.write(f"  Created:          {created}")
        self.stdout.write(f"  Updated:          {updated}")
        self.stdout.write(f"  Skipped (dupe):   {skipped_dupe}")
        self.stdout.write(f"  Skipped (invalid): {skipped_invalid}")

        if errors:
            self.stdout.write(self.style.WARNING(
                f"\n{len(errors)} problem(s) found:"
            ))
            for e in errors[:100]:   # cap at 100 to keep output readable
                self.stdout.write(f"  • {e}")
            if len(errors) > 100:
                self.stdout.write(f"  ... and {len(errors) - 100} more")

    @staticmethod
    def _clean(value):
        """Convert any cell value to a trimmed string, treating None/blank as empty."""
        if value is None:
            return ''
        return str(value).strip()